"""Contrôle d'un classeur « modèle standard » avant son utilisation par Orisflow.

Orisflow écrit dans le dernier classeur du dossier de référence (qui deviendra le modèle
standard : logo, volets figés, protections). Ce contrôle dit, avant la première génération :
- ce dont Orisflow a besoin pour fonctionner (onglet, libellés des lignes) — **bloquant** ;
- ce qu'Orisflow **ne conserverait pas** en réécrivant le classeur (simulé réellement :
  chargement puis enregistrement d'une copie, puis comparaison) — **bloquant** ;
- ce qui manque à la standardisation voulue (logo en haut à gauche, impossibilité d'insérer,
  supprimer ou déplacer des lignes et colonnes, structure verrouillée, volets figés, cellule
  de signature nommée) — **à corriger**.

Le fichier contrôlé n'est jamais modifié.
"""

from __future__ import annotations

import os
import re
import shutil
import tempfile
import zipfile
from typing import Any

from openpyxl import load_workbook

from .chemins import chemin_lecture
from .generation import FEUILLES_A_EXCLURE
from .reference_treso import NOM_FEUILLE_SYNTHESE
from .structure_modele import NOM_CELLULE_SIGNATURE, localiser_lignes

_REPERES_ENTETE_PIED = re.compile(rb'id="(LH|CH|RH|LF|CF|RF)"')
_FORMES = re.compile(rb"<xdr:(sp|grpSp|cxnSp)[ >]")


def _controle(libelle: str, ok: bool, detail: str, gravite: str) -> dict[str, Any]:
    """`gravite` : « bloquant » (Orisflow ne peut pas ou ne conserverait pas) ou « a_corriger »
    (standardisation incomplète). Un contrôle réussi garde sa gravité pour le tri."""
    return {"libelle": libelle, "ok": ok, "detail": detail, "gravite": gravite}


def _empreinte_fonctionnalites(chemin: str) -> dict[str, Any]:
    """Ce qu'on compare avant/après réécriture : tout ce que le standard doit garder."""
    classeur = load_workbook(chemin)
    feuille = next((classeur[n] for n in NOM_FEUILLE_SYNTHESE if n in classeur.sheetnames), None)
    if feuille is None:
        return {}
    protection = feuille.protection
    return {
        "images": len(feuille._images),
        "feuille_protegee": bool(protection.sheet),
        "mot_de_passe_feuille": bool(protection.password),
        "insertion_lignes_interdite": bool(protection.insertRows),
        "insertion_colonnes_interdite": bool(protection.insertColumns),
        "suppression_lignes_interdite": bool(protection.deleteRows),
        "suppression_colonnes_interdite": bool(protection.deleteColumns),
        "structure_verrouillee": bool(getattr(classeur.security, "lockStructure", False)),
        "volets_figes": feuille.freeze_panes,
        "noms_definis": sorted(classeur.defined_names.keys()),
        "cellules_fusionnees": sorted(str(r) for r in feuille.merged_cells.ranges),
        "largeur_colonne_a": feuille.column_dimensions["A"].width,
    }


def _formes_et_images_entete(chemin: str) -> tuple[int, bool]:
    """(nombre de formes/zones de texte, image placée dans l'en-tête ou le pied de page).
    Ni l'un ni l'autre ne survit à une réécriture par Orisflow."""
    formes = 0
    entete = False
    with zipfile.ZipFile(chemin) as archive:
        for nom in archive.namelist():
            if nom.startswith("xl/drawings/") and nom.endswith(".xml"):
                formes += len(_FORMES.findall(archive.read(nom)))
            elif nom.startswith("xl/drawings/vmlDrawing") and _REPERES_ENTETE_PIED.search(archive.read(nom)):
                entete = True
    return formes, entete


def _logo_en_haut_a_gauche(feuille) -> tuple[bool, str]:
    if not feuille._images:
        return False, "aucune image dans l'onglet « Synthèse » (Insertion > Images)"
    for image in feuille._images:
        ancre = image.anchor
        origine = getattr(ancre, "_from", None)
        if origine is not None and origine.col <= 1 and origine.row <= 2:
            return True, f"image ancrée en {chr(65 + origine.col)}{origine.row + 1}"
        position = getattr(ancre, "pos", None)  # ancre absolue (EMU)
        if position is not None and position.x < 1_500_000 and position.y < 1_000_000:
            return True, "image placée en haut à gauche"
    return False, "l'image n'est pas en haut à gauche (colonnes A-B, lignes 1 à 3)"


def controler_modele(chemin: str) -> dict[str, Any]:
    """Retourne {"ok", "bloquants", "a_corriger", "controles": [...]} ; `ok` = aucun bloquant."""
    lecture = chemin_lecture(chemin)
    controles: list[dict[str, Any]] = []

    try:
        with zipfile.ZipFile(lecture):
            pass
        classeur = load_workbook(lecture)
    except Exception as erreur:
        detail = (
            "le classeur est protégé par un mot de passe d'OUVERTURE (« Chiffrer avec mot de passe ») : Orisflow ne peut pas le lire. "
            "Utilisez seulement la protection de la feuille et du classeur."
            if "zip" in str(erreur).lower() or "not a zip" in str(erreur).lower()
            else f"{erreur.__class__.__name__}"
        )
        return {
            "ok": False, "bloquants": 1, "a_corriger": 0,
            "controles": [_controle("Le classeur est lisible", False, detail, "bloquant")],
        }

    feuille = next((classeur[n] for n in NOM_FEUILLE_SYNTHESE if n in classeur.sheetnames), None)
    controles.append(_controle("Onglet « Synthèse » présent", feuille is not None,
                               "trouvé" if feuille is not None else "introuvable", "bloquant"))
    if feuille is None:
        return _bilan(controles)

    lignes, erreurs = localiser_lignes(feuille)
    controles.append(_controle(
        "Lignes retrouvées par leur libellé (colonne A)", not erreurs,
        f"{len(lignes)} lignes reconnues" if not erreurs else " ; ".join(erreurs), "bloquant"))

    en_trop = [n for n in classeur.sheetnames if n in FEUILLES_A_EXCLURE]
    controles.append(_controle(
        "Seul l'onglet « Synthèse » est présent", not en_trop,
        "oui" if not en_trop else f"ces onglets seront retirés par Orisflow : {', '.join(en_trop)}", "a_corriger"))

    # Ce que la réécriture par Orisflow conserve : simulée, jamais supposée.
    dossier = tempfile.mkdtemp(prefix="orisflow_controle_")
    try:
        copie = os.path.join(dossier, "copie.xlsx")
        shutil.copyfile(lecture, os.path.join(dossier, "original.xlsx"))
        avant = _empreinte_fonctionnalites(os.path.join(dossier, "original.xlsx"))
        load_workbook(os.path.join(dossier, "original.xlsx")).save(copie)
        apres = _empreinte_fonctionnalites(copie)
    finally:
        shutil.rmtree(dossier, ignore_errors=True)
    perdus = [cle for cle in avant if avant[cle] != apres.get(cle)]
    formes, image_entete = _formes_et_images_entete(lecture)
    if formes:
        perdus.append(f"{formes} forme(s) ou zone(s) de texte")
    if image_entete:
        perdus.append("image placée dans l'en-tête/pied de page")
    controles.append(_controle(
        "Orisflow conserve tout ce que le modèle contient", not perdus,
        "logo, protections, volets figés, noms, largeurs : tout est conservé après réécriture"
        if not perdus else "serait perdu : " + ", ".join(perdus)
        + (" — remplacez les formes et images d'en-tête par une image insérée dans la feuille (Insertion > Images)"
           if formes or image_entete else ""),
        "bloquant"))

    ok_logo, detail_logo = _logo_en_haut_a_gauche(feuille)
    controles.append(_controle("Logo ORIS FINANCE en haut à gauche", ok_logo, detail_logo, "a_corriger"))

    protection = feuille.protection
    controles.append(_controle("Feuille « Synthèse » protégée", bool(protection.sheet),
                               "protégée" if protection.sheet else "Révision > Protéger la feuille", "a_corriger"))
    interdits = {
        "insérer des lignes": protection.insertRows, "insérer des colonnes": protection.insertColumns,
        "supprimer des lignes": protection.deleteRows, "supprimer des colonnes": protection.deleteColumns,
    }
    permis = [action for action, interdit in interdits.items() if not interdit]
    controles.append(_controle(
        "Impossible d'insérer ou supprimer lignes et colonnes", protection.sheet and not permis,
        "interdit" if protection.sheet and not permis
        else ("la feuille n'est pas protégée" if not protection.sheet
              else "encore permis : " + ", ".join(permis) + " (décochez ces cases dans « Protéger la feuille »)"),
        "a_corriger"))
    verrouille = bool(getattr(classeur.security, "lockStructure", False))
    controles.append(_controle("Structure du classeur verrouillée (onglets)", verrouille,
                               "verrouillée" if verrouille else "Révision > Protéger le classeur > Structure", "a_corriger"))
    controles.append(_controle("Volets figés (logo toujours visible)", feuille.freeze_panes is not None,
                               f"figés en {feuille.freeze_panes}" if feuille.freeze_panes else "Affichage > Figer les volets", "a_corriger"))

    nom = classeur.defined_names.get(NOM_CELLULE_SIGNATURE)
    cible = None
    if nom is not None:
        destinations = list(nom.destinations)
        cible = destinations[0] if destinations else None
    ok_nom = cible is not None and cible[0] in NOM_FEUILLE_SYNTHESE
    controles.append(_controle(
        f"Cellule de signature nommée « {NOM_CELLULE_SIGNATURE} »", ok_nom,
        f"pointe sur {cible[0]}!{cible[1]}" if ok_nom else "à créer : Formules > Gestionnaire de noms (voir standardisation.txt)",
        "a_corriger"))
    return _bilan(controles)


def _bilan(controles: list[dict[str, Any]]) -> dict[str, Any]:
    bloquants = sum(1 for c in controles if not c["ok"] and c["gravite"] == "bloquant")
    a_corriger = sum(1 for c in controles if not c["ok"] and c["gravite"] == "a_corriger")
    return {"ok": bloquants == 0, "bloquants": bloquants, "a_corriger": a_corriger, "controles": controles}
