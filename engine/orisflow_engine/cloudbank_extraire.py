"""Extraction d'un sous-ensemble d'un grand livre auxiliaire CloudBank (portage de `filtrer_grand_livre`)."""

from __future__ import annotations

import os
import re
from copy import copy
from typing import Any, Optional

import openpyxl

from .chemins import chemin_lecture
from .cloudbank_ecriture import enregistrer_classeur
from .cloudbank_referentiels import ErreurCloudBank


def filtrer_grand_livre(classeur_source, termes: list[str]):
    """Garde les sous-comptes dont le libellé contient un des `termes` (insensible à la casse, OU).

    Structure d'origine conservée (en-tête général, ligne d'agence, chapitre, en-tête de colonnes,
    sous-compte, transactions, TOTAL COMPTE), ainsi que largeurs, hauteurs, styles et fusions.
    Jamais de modification du classeur source (reconstruction dans un nouveau classeur).

    PIÈGE CONNU (09/10/2026, export ABCJTemp) : ce format déclare parfois une dimension de feuille
    fausse (`<dimension ref="A1"/>` pour des milliers de lignes). Conséquences vérifiées :
    - `load_workbook(..., read_only=True)` tronque la feuille à 1 ligne : TOUJOURS charger sans
      `read_only` (voir `extraire`) ;
    - `ws.delete_rows` corrompt silencieusement les cellules au-delà de la colonne B : d'où la
      reconstruction cellule par cellule ci-dessous plutôt qu'une suppression en place.
    """
    termes_n = [t.upper() for t in termes]
    src = classeur_source.active
    max_row = src.max_row
    max_col = src.max_column or 13

    def valeur_b(r: int):
        v = src.cell(r, 2).value
        return v.strip() if isinstance(v, str) else None

    conservees = set(range(1, min(7, max_row + 1)))  # en-tête général (6 premières lignes)
    retenus = []
    for r in range(7, max_row + 1):
        bs = valeur_b(r)
        if not bs or not re.match(r"^\d{13}\s*:", bs):
            continue
        if not any(terme in bs.upper() for terme in termes_n):
            continue
        retenus.append(bs)
        conservees.add(r)
        rr = r - 1
        if valeur_b(rr) == "Agence":
            conservees.add(rr)
            rr -= 1
        rr2 = rr
        while rr2 >= 7 and not valeur_b(rr2):
            rr2 -= 1
        if valeur_b(rr2) and re.match(r"^\d{5}\s*:", valeur_b(rr2)):
            conservees.add(rr2)
        rr3 = r
        while rr3 >= 7 and not (valeur_b(rr3) and valeur_b(rr3).startswith("Agence")
                                and ":" in valeur_b(rr3) and valeur_b(rr3) != "Agence"):
            rr3 -= 1
        if rr3 >= 7:
            conservees.add(rr3)
        rt = r + 1
        while rt <= max_row:
            fval = src.cell(rt, 6).value
            bval = valeur_b(rt)
            if bval and (re.match(r"^\d{13}\s*:", bval) or re.match(r"^\d{5}\s*:", bval)
                         or (bval.startswith("Agence") and bval != "Agence")):
                break
            conservees.add(rt)
            if fval == "TOTAL COMPTE":
                break
            rt += 1

    dst_wb = openpyxl.Workbook()
    dst = dst_wb.active
    dst.title = src.title
    for c in range(1, max_col + 1):
        lettre = src.cell(1, c).column_letter
        dim = src.column_dimensions.get(lettre)
        if dim is not None and dim.width:
            dst.column_dimensions[lettre].width = dim.width

    index_dest = {}
    for dst_r, src_r in enumerate(sorted(conservees), start=1):
        index_dest[src_r] = dst_r
        hauteur = src.row_dimensions[src_r].height
        if hauteur is not None:
            dst.row_dimensions[dst_r].height = hauteur
        for c in range(1, max_col + 1):
            sc = src.cell(src_r, c)
            dc = dst.cell(dst_r, c, sc.value)
            if sc.has_style:
                dc.font, dc.fill, dc.border, dc.alignment = copy(sc.font), copy(sc.fill), copy(sc.border), copy(sc.alignment)
                dc.number_format = sc.number_format
    for rng in src.merged_cells.ranges:
        if rng.min_row in index_dest and rng.max_row in index_dest:
            dst.merge_cells(start_row=index_dest[rng.min_row], start_column=rng.min_col,
                            end_row=index_dest[rng.max_row], end_column=rng.max_col)
    total = sum(dst.cell(r, 8).value or 0 for r in range(1, dst.max_row + 1) if dst.cell(r, 6).value == "TOTAL COMPTE")
    return dst_wb, retenus, total


def extraire(chemin_grand_livre: str, sous_comptes: list[str], dossier_sortie: str,
             nom: Optional[str] = None) -> dict[str, Any]:
    termes = [t.strip() for t in sous_comptes if t and t.strip()]
    if not termes:
        raise ErreurCloudBank("Extraire : indiquez au moins un terme à chercher dans les sous-comptes.")
    # Jamais read_only (voir le piège ci-dessus) ; jamais sauvegardé : le fichier source reste intact.
    source = openpyxl.load_workbook(chemin_lecture(chemin_grand_livre))
    destination, retenus, total = filtrer_grand_livre(source, termes)
    base = os.path.splitext(os.path.basename(chemin_grand_livre))[0]
    chemin = enregistrer_classeur(destination, dossier_sortie, nom or f"{base} - extrait.xlsx")
    return {"fichier": chemin, "termes": termes, "sous_comptes_retenus": retenus,
            "nombre_sous_comptes": len(retenus), "lignes": destination.active.max_row, "somme_total_compte": total}
