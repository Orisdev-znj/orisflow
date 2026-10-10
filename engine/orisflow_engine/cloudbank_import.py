"""Lecture des fichiers échangés avec Claude (format d'échange du §6 du document de conception).

- `MAPPING A VALIDER - <AGENCE> - <Mois Année>.xlsx` : une ligne par dépense, colonnes
  LIBELLE | MONTANT | COMPTE PROPOSE | INTITULE | STATUT | SOURCE | COMMENTAIRE ;
- fichier déjà final (extourne, extrait de grand livre) : simplement résumé, puis déposé.

Les fichiers importés ne sont jamais modifiés.
"""

from __future__ import annotations

import os
import re
import shutil
from typing import Any

import openpyxl

from .chemins import chemin_lecture
from .cloudbank_ecriture import chemin_sans_ecrasement
from .cloudbank_mapping import chercher_pcemf, mapper_ligne
from .cloudbank_referentiels import ErreurCloudBank, Referentiels, normaliser
from .cloudbank_resolution_agence import resoudre_compte_agence

COLONNES_OBLIGATOIRES = {"libelle": "LIBELLE", "montant": "MONTANT", "compte propose": "COMPTE PROPOSE", "statut": "STATUT"}
STATUTS = {"valide": "valide", "a confirmer": "a_confirmer", "non trouve": "non_trouve"}
_NOM_MAPPING = re.compile(r"^MAPPING A VALIDER\s*-\s*(?P<agence>.+?)\s*-\s*(?P<mois>[A-Za-zÀ-ÿ]+)\s+(?P<annee>\d{4})", re.IGNORECASE)


def _nombre(valeur):
    if valeur is None or valeur == "":
        return None
    if isinstance(valeur, (int, float)):
        return valeur
    texte = re.sub(r"[\s ]", "", str(valeur)).replace(",", ".")
    try:
        nombre = float(texte)
    except ValueError:
        return None
    return int(nombre) if nombre.is_integer() else nombre


def lire_nom_mapping(ref: Referentiels, nom_fichier: str) -> dict[str, Any]:
    """Agence et période d'après le nom du fichier (convention du §6)."""
    m = _NOM_MAPPING.match(os.path.basename(nom_fichier))
    if not m:
        return {"conforme": False, "agence": None, "agence_nom": None, "mois": None, "annee": None}
    code = ref.code_agence(m.group("agence"))
    return {"conforme": True, "agence": code, "agence_nom": ref.agences.get(code) if code else m.group("agence"),
            "mois": m.group("mois").capitalize(), "annee": m.group("annee")}


def lire_fichier_mapping(ref: Referentiels, chemin: str) -> dict[str, Any]:
    """Lignes du fichier « MAPPING A VALIDER », avec compte réel résolu PAR AGENCE (règle n°7)."""
    try:
        classeur = openpyxl.load_workbook(chemin_lecture(chemin), data_only=True)
    except Exception:
        raise ErreurCloudBank("Ce fichier n'a pas pu être lu : est-il bien un fichier Excel (.xlsx) ?") from None
    feuille = classeur.worksheets[0]

    ligne_entete, colonnes = None, {}
    for r in range(1, min(feuille.max_row, 15) + 1):
        trouvees = {normaliser(feuille.cell(r, c).value): c for c in range(1, feuille.max_column + 1) if feuille.cell(r, c).value}
        if "libelle" in trouvees and "montant" in trouvees:
            ligne_entete, colonnes = r, trouvees
            break
    manquantes = [nom for cle, nom in COLONNES_OBLIGATOIRES.items() if cle not in colonnes]
    if ligne_entete is None or manquantes:
        raise ErreurCloudBank("Ce fichier ne correspond pas au format attendu : les colonnes "
                              + ", ".join(manquantes or list(COLONNES_OBLIGATOIRES.values())) + " sont introuvables.")

    contexte = lire_nom_mapping(ref, chemin)
    lignes, total = [], 0
    for r in range(ligne_entete + 1, feuille.max_row + 1):
        libelle = feuille.cell(r, colonnes["libelle"]).value
        montant = _nombre(feuille.cell(r, colonnes["montant"]).value)
        if (libelle is None or str(libelle).strip() == "") and montant is None:
            continue
        libelle = str(libelle or "").strip()
        compte_brut = feuille.cell(r, colonnes["compte propose"]).value
        compte = re.sub(r"\D", "", str(compte_brut)) if compte_brut not in (None, "") else None
        statut = STATUTS.get(normaliser(feuille.cell(r, colonnes["statut"]).value))
        intitule = feuille.cell(r, colonnes["intitule"]).value if "intitule" in colonnes else None
        source = feuille.cell(r, colonnes["source"]).value if "source" in colonnes else None
        commentaire = feuille.cell(r, colonnes["commentaire"]).value if "commentaire" in colonnes else None
        avertissement = None

        if statut is None:  # statut absent ou inconnu : le moteur recalcule, sans rien présumer
            proposition = mapper_ligne(ref, libelle)
            statut, compte, intitule, source = proposition["statut"], proposition["compte"], proposition["intitule"], proposition["source"]
            avertissement = "Le statut n'était pas indiqué : il a été recalculé par Orisflow."
        if compte is not None and not re.fullmatch(r"\d{13}", compte):
            avertissement = f"Numéro de compte non conforme (13 chiffres attendus) : {compte_brut}"
            compte, statut = None, "non_trouve"
        if statut == "non_trouve":
            compte = None
        elif compte and not intitule:
            intitule = ref.cloudbank.get(compte, {}).get("intitule")

        compte_reel, avert_agence = (compte, None)
        if compte and contexte["agence"]:
            compte_reel, avert_agence = resoudre_compte_agence(ref, compte, contexte["agence"])
        if montant is None:
            avertissement = "Montant illisible ou absent."
        total += montant or 0
        lignes.append({
            "id": len(lignes), "libelle": libelle, "montant": montant, "compte": compte, "compte_reel": compte_reel,
            "intitule": intitule, "statut": statut, "source": source, "commentaire": commentaire,
            "sens": "credit" if commentaire and "credit" in normaliser(commentaire).split() else "debit",
            "avertissement": avertissement, "avertissement_agence": avert_agence,
            "suggestions": chercher_pcemf(ref, libelle) if statut == "non_trouve" else None,
        })
    if not lignes:
        raise ErreurCloudBank("Ce fichier ne contient aucune ligne de dépense à confirmer.")
    resume = {s: sum(1 for l in lignes if l["statut"] == s) for s in ("valide", "a_confirmer", "non_trouve")}
    return {"contexte": contexte, "lignes": lignes, "resume": resume, "total": total}


def resumer_fichier_resultat(chemin: str) -> dict[str, Any]:
    """Résumé d'un fichier déjà final (extourne, extrait) : feuilles, lignes, totaux. Lecture seule."""
    try:
        classeur = openpyxl.load_workbook(chemin_lecture(chemin), data_only=True)  # jamais read_only (voir cloudbank_extraire)
    except Exception:
        raise ErreurCloudBank("Ce fichier n'a pas pu être lu : est-il bien un fichier Excel (.xlsx) ?") from None
    feuilles, avertissements = [], []
    for feuille in classeur.worksheets:
        lignes_compte, debit, credit, total_comptes = 0, 0, 0, 0
        for r in range(2, feuille.max_row + 1):
            a = feuille.cell(r, 1).value
            if a is not None and re.fullmatch(r"\d{13}", str(a).strip()):
                lignes_compte += 1
                debit += _nombre(feuille.cell(r, 4).value) or 0
                credit += _nombre(feuille.cell(r, 5).value) or 0
            if feuille.cell(r, 6).value == "TOTAL COMPTE":
                total_comptes += _nombre(feuille.cell(r, 8).value) or 0
        feuilles.append({"feuille": feuille.title, "lignes": feuille.max_row, "lignes_ecriture": lignes_compte,
                         "total_debit": debit, "total_credit": credit, "ecart": debit - credit,
                         "somme_total_compte": total_comptes})
        if lignes_compte and debit != credit:
            avertissements.append(f"Feuille « {feuille.title} » : débit et crédit différents (écart {debit - credit}).")
    return {"feuilles": feuilles, "avertissements": avertissements}


def deposer_fichier(chemin: str, dossier_sortie: str) -> dict[str, Any]:
    """Copie un fichier déjà final dans le dossier de résultats, sans jamais écraser un fichier existant."""
    resume = resumer_fichier_resultat(chemin)
    os.makedirs(chemin_lecture(dossier_sortie), exist_ok=True)
    cible = chemin_sans_ecrasement(dossier_sortie, os.path.basename(chemin))
    shutil.copyfile(chemin_lecture(chemin), chemin_lecture(cible))
    return {"fichier": cible, **resume}
