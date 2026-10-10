"""Lecture SEULE du dernier classeur de trésorerie existant, pour comparer un
comptage du jour à celui de la veille. N'écrit jamais dans ce classeur.

La recherche du fichier le plus récent reprend `trouver_tresorerie_recente` du
script B (date la plus récente lue dans le nom du fichier, sous-dossiers inclus,
fichiers `~$` ignorés).
"""

from __future__ import annotations

import glob
import os
import re
from datetime import date
from typing import Optional

from openpyxl import load_workbook

from .recalcul import recalculer_classeur
from .regles_agences import AGENCE_COLONNE
from .regles_banques import BONS_DE_CAISSE, CELLULE_CHAMP_MANUEL
from .structure_modele import localiser_lignes
from .chemins import chemin_lecture

NOM_FEUILLE_SYNTHESE = ("Synthèse", "Synthese")


def extraire_date_nom(chemin: str) -> date:
    nom = os.path.basename(chemin)
    correspondance = re.search(r"(\d{2})[-\s_](\d{2})[-\s_](\d{4})", nom)
    if correspondance:
        try:
            return date(int(correspondance.group(3)), int(correspondance.group(2)), int(correspondance.group(1)))
        except ValueError:
            pass
    return date.min


def trouver_classeur_recent(dossier: str, avant: Optional[date] = None) -> Optional[str]:
    """`avant` (optionnel) : n'accepte que les classeurs dont la date lue dans le nom est
    strictement antérieure à `avant`. Nécessaire pour la génération (voir `generation.py`) :
    sans ce filtre, si le dossier de référence contient déjà un classeur daté du jour qu'on
    s'apprête à générer (ou d'un jour postérieur — rattrapage, saisie en avance...), ce
    classeur serait pris comme son propre modèle, et son propre J-1 comme référence de
    lui-même. Bug corrigé le 01/10/2026 (CLAUDE.md), repéré après l'ajout d'un classeur du
    30/09 dans le dossier de référence le jour même où Orisflow générait pour le 30/09."""
    if not dossier or not os.path.isdir(dossier):
        return None
    fichiers = glob.glob(os.path.join(dossier, "**", "*.xlsx"), recursive=True)
    fichiers += glob.glob(os.path.join(dossier, "**", "*.xlsm"), recursive=True)
    fichiers = [f for f in fichiers if not os.path.basename(f).startswith("~$")]
    if avant is not None:
        fichiers = [f for f in fichiers if extraire_date_nom(f) < avant]
    if not fichiers:
        return None
    return max(fichiers, key=extraire_date_nom)


def _lire_valeurs(chemin: str) -> tuple[dict[str, int], dict[str, int]]:
    """(totaux de comptes par agence, valeurs des champs manuels) lus dans les valeurs
    calculées du classeur. Ce qui est absent ou illisible est simplement omis."""
    totaux: dict[str, int] = {}
    manuels: dict[str, int] = {}
    try:
        classeur = load_workbook(chemin_lecture(chemin), data_only=True, read_only=True)
        try:
            feuille = next((classeur[n] for n in NOM_FEUILLE_SYNTHESE if n in classeur.sheetnames), None)
            if feuille is None:
                return totaux, manuels
            cles = ["total_comptes"] + sorted({ligne for ligne, _ in CELLULE_CHAMP_MANUEL.values()})
            lignes, _erreurs = localiser_lignes(feuille, cles)  # chaque ligne est utilisée indépendamment
            if "total_comptes" in lignes:
                for agence_cle, colonne in AGENCE_COLONNE.items():
                    valeur = feuille[f"{colonne}{lignes['total_comptes']}"].value
                    if isinstance(valeur, (int, float)):
                        totaux[agence_cle] = int(valeur)
            for champ, (cle_ligne, agence) in CELLULE_CHAMP_MANUEL.items():
                if cle_ligne not in lignes:
                    continue
                valeur = feuille[f"{AGENCE_COLONNE[agence]}{lignes[cle_ligne]}"].value
                if not isinstance(valeur, (int, float)):
                    continue
                if champ == "uba_solde_banque":
                    valeur -= sum(BONS_DE_CAISSE.get(("uba", "akwa"), []))
                manuels[champ] = int(valeur)
        finally:
            classeur.close()
    except Exception:
        # Un classeur illisible ne doit jamais empêcher l'import : on renvoie ce qui a pu être lu.
        pass
    return totaux, manuels


def lire_totaux_comptes_precedents(dossier_reference: str) -> dict:
    """Lit, en LECTURE SEULE, le dernier classeur du dossier de référence.

    Résultat : {"chemin", "date", "totaux": {agence: total de comptes},
    "valeurs_manuelles": {champ manuel: valeur de la veille}}.
    Si le classeur n'a jamais été recalculé par Excel (aucun total lisible), il est recalculé
    via LibreOffice (copie en cache, voir `recalcul.py`) avant d'abandonner.
    """
    chemin = trouver_classeur_recent(dossier_reference)
    if chemin is None:
        return {"chemin": None, "date": None, "totaux": {}, "valeurs_manuelles": {}}

    totaux, manuels = _lire_valeurs(chemin)
    if not totaux:
        chemin_recalcule = recalculer_classeur(chemin_lecture(chemin))
        if chemin_recalcule is not None:
            totaux, manuels = _lire_valeurs(chemin_recalcule)

    jour = extraire_date_nom(chemin)
    return {
        "chemin": chemin,
        "date": jour.isoformat() if jour != date.min else None,
        "totaux": totaux,
        "valeurs_manuelles": manuels,
    }
