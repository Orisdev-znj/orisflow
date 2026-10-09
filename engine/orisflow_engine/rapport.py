"""Export du rapport d'analyse en classeur Excel (demande du 10/10/2026) : permet de garder
une trace lisible hors de l'application, ou de la transmettre pour un contrôle a posteriori,
sans dépendre du presse-papiers (voir aussi EcranResultats.tsx, « Copier le rapport
d'analyse »). N'écrase jamais un fichier existant à l'insu de l'utilisateur : le nom et
l'emplacement sont choisis par lui, via la boîte de dialogue « Enregistrer sous » côté
Electron — ce module se contente d'écrire à l'endroit indiqué."""

from __future__ import annotations

import os
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

COLONNES = ["Fichier", "Type détecté", "Agence", "Confiance agence", "Comptes", "Niveau", "Messages"]

NIVEAU_LIBELLES = {"information": "Conforme", "avertissement": "À vérifier", "bloquant": "Rejeté"}


def exporter_rapport_excel(
    fichiers: list[dict[str, Any]],
    chemin: str,
    journal_etapes: Optional[list[str]] = None,
) -> None:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Rapport d'analyse"

    ligne = 1
    if journal_etapes:
        feuille.cell(row=ligne, column=1, value="Étapes de l'analyse").font = Font(bold=True)
        ligne += 1
        for etape in journal_etapes:
            feuille.cell(row=ligne, column=1, value=f"- {etape}")
            ligne += 1
        ligne += 1  # ligne vide avant le tableau

    ligne_entete = ligne
    for colonne, titre in enumerate(COLONNES, start=1):
        feuille.cell(row=ligne_entete, column=colonne, value=titre).font = Font(bold=True)
    ligne += 1

    for f in fichiers:
        agence = "—" if f.get("confiance_agence") == "sans_objet" else (f.get("agence_libelle") or "Non reconnue")
        valeurs = [
            f.get("nom"),
            f.get("type_libelle") or "—",
            agence,
            f.get("confiance_agence"),
            f.get("total_comptes"),
            NIVEAU_LIBELLES.get(f.get("niveau"), f.get("niveau")),
            "\n".join(f.get("messages") or []),
        ]
        for colonne, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
            if colonne == len(COLONNES):
                cellule.alignment = Alignment(wrap_text=True, vertical="top")
        ligne += 1

    for colonne_lettre, largeur in zip("ABCDEFG", [40, 26, 16, 20, 10, 12, 90]):
        feuille.column_dimensions[colonne_lettre].width = largeur
    feuille.freeze_panes = f"A{ligne_entete + 1}"

    os.makedirs(os.path.dirname(chemin) or ".", exist_ok=True)
    classeur.save(chemin)
