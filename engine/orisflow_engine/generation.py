"""Génération du classeur de trésorerie journalière.

Principes (CLAUDE.md) :
- On part toujours du **dernier classeur existant** du dossier de référence, jamais modifié
  ni écrasé ; le résultat est un nouveau fichier.
- Le classeur porte par défaut la date de **la veille** (la trésorerie traitée un matin
  concerne la journée précédente, confirmé le 01/10/2026).
- Chaque ligne est retrouvée par son **libellé** (`structure_modele`) : si le modèle n'a pas
  la structure attendue, rien n'est écrit.
- Les « J-1 » sont lus dans le modèle **avant** toute écriture, sur ses valeurs calculées
  (recalculées via LibreOffice si le modèle n'a jamais été rouvert dans Excel).
- Un fichier bloquant n'est jamais utilisé ; un doublon est ignoré sans être une anomalie.
- « Suivi de la treso » et les feuilles mortes de l'ancienne méthode sont retirés.
"""

from __future__ import annotations

import os
from datetime import date, timedelta
from typing import Any, Callable, Optional

from openpyxl import load_workbook

from . import carnet
from .classification import classer_fichiers
from .generation_banques import ecrire_banques, releves_fictifs, soldes_du_jour
from .generation_caisses import ecrire_caisses
from .recalcul import recalculer_classeur
from .reference_treso import NOM_FEUILLE_SYNTHESE, trouver_classeur_recent
from .regles_agences import AGENCE_COLONNE, AGENCE_LIBELLES
from .structure_modele import LIGNES_BANQUES, lignes_comptes, localiser_lignes, valeur_numerique
from .chemins import chemin_lecture

FEUILLES_A_EXCLURE = (
    "Suivi de la treso",
    "Agences",
    "Dépôts",
    "Caisse",
    "SYNTHESE 2025",
    "Feuil1",
)

# Catégorie de comptes (comptes.py) -> clé de ligne (structure_modele). Collectes et salariés
# n'ont aucune règle : remis à 0, comme le fait le script d'origine.
LIGNE_PAR_CATEGORIE = {
    "Courants": "courants",
    "Cheques": "cheques",
    "Epargne": "epargne",
    "Garanties": "garanties",
    "Fonctionnaires": "fonctionnaires",
}

# Saisies manuelles enregistrées au carnet. Maviance n'y figure pas (décision du 05/10/2026).
_NOMS_CARNET_MANUELS = {
    "uv_orange": "uv:orange_money",
    "uv_mtn": "uv:mtn_momo",
    "ecobank": "ecobank:akwa",
    "access_bank": "access_bank:marche_central",
    "uba_solde_banque": "uba:akwa",
}


def _nom_fichier_du_jour(jour: date) -> str:
    # Deux espaces avant la date : convention de tous les classeurs réels.
    return f"TRESORERIE JOURNALIÈRE et TDB DU  {jour.day:02d} {jour.month:02d} {jour.year}.xlsx"


def _chemin_disponible(dossier: str, nom: str) -> str:
    """Ajoute un suffixe numéroté si `nom` existe déjà dans `dossier` (jamais d'écrasement)."""
    base, extension = os.path.splitext(nom)
    chemin = os.path.join(dossier, nom)
    compteur = 2
    while os.path.exists(chemin_lecture(chemin)):
        chemin = os.path.join(dossier, f"{base} ({compteur}){extension}")
        compteur += 1
    return chemin


def _lignes_j1(lignes: dict[str, int]) -> list[int]:
    return list(lignes_comptes(lignes)) + [lignes["depots"], lignes["engagements"]] + [lignes[c] for c in LIGNES_BANQUES]


def _cellules_sans_valeur(feuille, feuille_valeurs, numeros_lignes: list[int]) -> int:
    """Cellules-formules du modèle sans valeur calculée (classeur jamais recalculé)."""
    manquantes = 0
    for colonne in AGENCE_COLONNE.values():
        for r in numeros_lignes:
            brute = feuille[f"{colonne}{r}"].value
            if isinstance(brute, str) and brute.startswith("=") and feuille_valeurs[f"{colonne}{r}"].value is None:
                manquantes += 1
    return manquantes


def _feuille_synthese(classeur):
    return next((classeur[n] for n in NOM_FEUILLE_SYNTHESE if n in classeur.sheetnames), None)


def generer_classeur(
    fichiers: list[str],
    dossier_reference: str,
    dossier_sortie: str,
    jour: Optional[date] = None,
    valeurs_manuelles: Optional[dict[str, Any]] = None,
    dossier_carnet: Optional[str] = None,
    releves_saisis: Optional[dict[str, int]] = None,
    table_comptes: Optional[dict[str, list[str]]] = None,
    dossier_cache: Optional[str] = None,
    sur_etape: Optional[Callable[[str], None]] = None,
) -> dict[str, Any]:
    etape = sur_etape or (lambda _message: None)
    jour = jour or (date.today() - timedelta(days=1))
    valeurs_manuelles = valeurs_manuelles or {}

    etape("Lecture des fichiers importés…")
    classement = classer_fichiers(
        fichiers, dossier_reference, dossier_carnet=dossier_carnet, jour=jour,
        table_comptes=table_comptes, dossier_cache=dossier_cache,
    )
    # `avant=jour` : un classeur daté du jour généré (ou plus tard) n'est jamais son propre modèle.
    chemin_modele = trouver_classeur_recent(dossier_reference, avant=jour)
    if chemin_modele is None:
        return {
            "ok": False,
            "erreur": "Aucun classeur de référence trouvé : impossible de savoir de quel modèle partir.",
            "classement": classement,
        }

    etape("Lecture du classeur de référence…")
    classeur = load_workbook(chemin_lecture(chemin_modele))  # formules et mise en forme conservées
    feuille = _feuille_synthese(classeur)
    if feuille is None:
        return {
            "ok": False,
            "erreur": "La feuille « Synthèse » est introuvable dans le classeur de référence.",
            "classement": classement,
        }
    lignes, erreurs_structure = localiser_lignes(feuille)
    if erreurs_structure:
        return {
            "ok": False,
            "erreur": (
                f"Le classeur de référence « {os.path.basename(chemin_modele)} » n'a pas la structure attendue : "
                + " ; ".join(erreurs_structure)
                + ". Rien n'a été généré : vérifiez les libellés de la colonne A de la feuille « Synthèse »."
            ),
            "classement": classement,
        }
    for nom_feuille in FEUILLES_A_EXCLURE:
        if nom_feuille in classeur.sheetnames:
            del classeur[nom_feuille]

    feuille_valeurs = _feuille_synthese(load_workbook(chemin_lecture(chemin_modele), data_only=True))
    if _cellules_sans_valeur(feuille, feuille_valeurs, _lignes_j1(lignes)):
        etape("Recalcul du classeur de référence (LibreOffice)…")
        chemin_recalcule = recalculer_classeur(chemin_lecture(chemin_modele))
        if chemin_recalcule is not None:
            recalcule = _feuille_synthese(load_workbook(chemin_recalcule, data_only=True))
            if recalcule is not None:
                feuille_valeurs = recalcule
    manquantes = _cellules_sans_valeur(feuille, feuille_valeurs, _lignes_j1(lignes))
    avertissements_modele = (
        [
            "Le classeur de référence n'a pas été recalculé par Excel depuis son enregistrement, et LibreOffice "
            f"n'a pas pu le faire : {manquantes} cellule(s) utilisée(s) pour le « J-1 » ont été comptées comme 0. "
            "Ouvrez ce classeur dans Excel, enregistrez-le, puis régénérez pour un J-1 fiable."
        ]
        if manquantes
        else []
    )

    # J-1 : valeurs du modèle, lues AVANT toute écriture (sinon on relirait nos propres valeurs).
    etape("Écriture des comptes, dépôts et engagements…")
    for colonne in AGENCE_COLONNE.values():
        feuille[f"{colonne}{lignes['total_comptes_j1']}"] = sum(
            valeur_numerique(feuille_valeurs, f"{colonne}{r}") for r in lignes_comptes(lignes)
        )
        feuille[f"{colonne}{lignes['depots_j1']}"] = valeur_numerique(feuille_valeurs, f"{colonne}{lignes['depots']}")
        feuille[f"{colonne}{lignes['engagements_j1']}"] = valeur_numerique(
            feuille_valeurs, f"{colonne}{lignes['engagements']}"
        )

    agences_comptes: list[str] = []
    agences_balance: list[str] = []
    fichiers_ignores: list[str] = []
    doublons_ignores: list[str] = []
    recap: dict[str, dict[str, Any]] = {}
    for f in classement["fichiers"]:
        if f["type_detecte"] not in ("compte", "balance_classe3"):
            continue
        if f["niveau"] == "bloquant" or f["agence_detectee"] is None:
            (doublons_ignores if f.get("est_doublon") else fichiers_ignores).append(f["nom"])
            continue
        agence = f["agence_detectee"]
        colonne = AGENCE_COLONNE[agence]
        ligne_recap = recap.setdefault(agence, {"agence": AGENCE_LIBELLES[agence]})

        if f["type_detecte"] == "compte":
            comptages = f["comptages"]
            if comptages is None:
                fichiers_ignores.append(f["nom"])
                continue
            for categorie, cle_ligne in LIGNE_PAR_CATEGORIE.items():
                feuille[f"{colonne}{lignes[cle_ligne]}"] = comptages.get(categorie, 0)
            feuille[f"{colonne}{lignes['collectes']}"] = 0
            feuille[f"{colonne}{lignes['salaries']}"] = 0
            ligne_recap["comptes"] = sum(comptages.get(c, 0) for c in LIGNE_PAR_CATEGORIE)
            agences_comptes.append(agence)
        else:
            if f["depots"] is None or f["engagements"] is None:
                fichiers_ignores.append(f["nom"])
                continue
            feuille[f"{colonne}{lignes['depots']}"] = f["depots"]
            feuille[f"{colonne}{lignes['engagements']}"] = f["engagements"]
            ligne_recap["depots"] = f["depots"]
            ligne_recap["engagements"] = f["engagements"]
            agences_balance.append(agence)

    # Relevés absents : valeur saisie, ou valeur de la veille (carnet) quand l'utilisateur l'a
    # choisie — injectés comme des relevés lus, jamais consignés au carnet du jour.
    saisis = {cle: int(v) for cle, v in (releves_saisis or {}).items() if v is not None}
    injectes = releves_fictifs(saisis, "saisi")
    absents_avec_veille = {
        m["cle"]: m["veille"]
        for m in classement["releves_manquants"]
        if m["cle"] not in saisis and m["veille"] is not None
    }
    injectes += releves_fictifs(absents_avec_veille, "veille")

    etape("Écriture des banques et des caisses…")
    banques = ecrire_banques(
        feuille, feuille_valeurs, lignes, {**classement, "fichiers": classement["fichiers"] + injectes}, valeurs_manuelles
    )
    caisses = ecrire_caisses(feuille, feuille_valeurs, classement)
    for agence, montant in caisses["montants"].items():
        recap.setdefault(agence, {"agence": AGENCE_LIBELLES[agence]})["caisse"] = montant
    fichiers_ignores.extend(banques["fichiers_ignores"])
    doublons_ignores.extend(banques["doublons_ignores"])

    etape("Enregistrement du classeur…")
    os.makedirs(chemin_lecture(dossier_sortie), exist_ok=True)
    chemin_sortie = _chemin_disponible(dossier_sortie, _nom_fichier_du_jour(jour))
    classeur.save(chemin_lecture(chemin_sortie))

    # Carnet : soldes réellement lus ou saisis aujourd'hui (jamais ceux repris de la veille).
    if dossier_carnet:
        soldes = soldes_du_jour(classement["fichiers"] + [i for i in injectes if i["origine"] == "saisi"])
        for champ, cle_carnet in _NOMS_CARNET_MANUELS.items():
            if valeurs_manuelles.get(champ) is not None:
                soldes[cle_carnet] = int(valeurs_manuelles[champ])
        carnet.enregistrer(os.path.join(dossier_carnet, carnet.NOM_FICHIER), jour, soldes)

    ordre = list(AGENCE_COLONNE)
    return {
        "ok": True,
        "chemin_genere": chemin_sortie,
        "date": jour.isoformat(),
        "modele_utilise": chemin_modele,
        "agences_mises_a_jour": [AGENCE_LIBELLES[c] for c in agences_comptes],
        "agences_non_mises_a_jour": [AGENCE_LIBELLES[c] for c in AGENCE_COLONNE if c not in agences_comptes],
        "agences_balance_mises_a_jour": [AGENCE_LIBELLES[c] for c in agences_balance],
        "agences_banques_mises_a_jour": [AGENCE_LIBELLES[c] for c in banques["agences"]],
        "agences_caisses_mises_a_jour": [AGENCE_LIBELLES[c] for c in caisses["agences"]],
        "avertissements_banques": banques["avertissements"] + caisses["avertissements"],
        "avertissements_modele": avertissements_modele,
        "releves_repris_de_la_veille": [i["nom"] for i in injectes if i["origine"] == "veille"],
        "releves_saisis": [i["nom"] for i in injectes if i["origine"] == "saisi"],
        "fichiers_ignores": fichiers_ignores,
        "doublons_ignores": doublons_ignores,
        "recapitulatif_agences": [recap[a] for a in sorted(recap, key=ordre.index)],
        "recapitulatif_banques": banques["recapitulatif"],
        "classement": classement,
    }
