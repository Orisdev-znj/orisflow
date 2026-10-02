"""Tests du sprint 4 : génération d'un nouveau classeur de trésorerie."""

import os
import sys
from datetime import date, timedelta

import openpyxl
import pymupdf
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from orisflow_engine.generation import generer_classeur

# Positions des colonnes observées sur les vrais documents CloudBank (voir CLAUDE.md §19.1
# et tests/test_balance_pdf.py).
_ANCRES = [231, 291, 351, 411, 471, 531]
_DECALAGE = 21


def _extraction_comptes(chemin, prefixes):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    for _ in range(23):
        feuille.append([])
    feuille.append(["N°", "Numero de compte"])
    for i, prefixe in enumerate(prefixes, start=1):
        feuille.append([i, f"{prefixe}-{i:06d}-00"])
    classeur.save(chemin)


def _extraction_balance_classe3(chemin, groupe, depots, engagements):
    """PDF synthétique de balance CloudBank classe 3 (dépôts/engagements), même mise en
    page que les vrais documents (voir CLAUDE.md §19.1 et test_balance_pdf.py)."""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((10, 50), "Balance generale consolidée   Chapitre de : 3  à : 3")
    page.insert_text((10, 70), f"Groupe: {groupe}   {groupe}")
    page.insert_text((10, 150), "Compte")
    page.insert_text((71, 150), "Intitulé")
    for i, x in enumerate(_ANCRES):
        page.insert_text((x, 150), "Debit" if i % 2 == 0 else "Crédit")
    # Valeurs : [Débit début, Crédit début, Débit MVT, Crédit MVT, Débit fin (engagements), Crédit fin (dépôts)]
    page.insert_text((10, 175), "Total Classe : 3 INTITULE TEST")
    for x, valeur in zip(_ANCRES, [0, 0, 0, 0, engagements, depots]):
        page.insert_text((x + _DECALAGE, 175), str(valeur))
    document.save(chemin)
    document.close()


def _classeur_modele(dossier, nom_fichier="TRESORERIE JOURNALIÈRE et TDB DU  10 09 2026.xlsx"):
    """Classeur de référence minimal, avec une feuille « Suivi de la treso » à exclure."""
    classeur = openpyxl.Workbook()
    synthese = classeur.active
    synthese.title = "Synthèse"
    synthese["A7"] = "COMPTES  COURANT ENTREPRISES"
    synthese["C7"] = 50  # ancienne valeur Akwa, doit se retrouver en C17 (J-1)
    synthese["C16"] = "=SUM(C7:C15)"  # formule à préserver
    synthese["F7"] = 30  # ancienne valeur Bafoussam
    synthese["F16"] = "=SUM(F7:F15)"
    synthese["A20"] = "ENCOURS  DEPOTS"
    synthese["C20"] = 1_000_000  # ancien dépôt Akwa, doit se retrouver en C21 (J-1)
    synthese["C22"] = "=C20-C21"  # formule à préserver
    synthese["F20"] = 2_000_000  # ancien dépôt Bafoussam
    synthese["A23"] = "ENCOURS ENGAGEMENTS"
    synthese["C23"] = 500_000  # ancien engagement Akwa, doit se retrouver en C24 (J-1)
    synthese["C25"] = "=C23-C24"  # formule à préserver
    synthese["F23"] = 800_000  # ancien engagement Bafoussam

    suivi = classeur.create_sheet("Suivi de la treso")
    suivi["A1"] = "Ne doit jamais apparaître dans le fichier généré"

    # Feuilles mortes héritées de l'ancienne méthode (voir CLAUDE.md, 29/09/2026) :
    # aucune formule de Synthèse ne les référence, elles doivent disparaître elles aussi.
    for nom in ("Agences", "Dépôts", "Caisse", "SYNTHESE 2025", "Feuil1"):
        classeur.create_sheet(nom)["A1"] = "Feuille morte, à exclure"

    chemin = dossier / nom_fichier
    classeur.save(chemin)
    return str(chemin)


@pytest.fixture
def contexte(tmp_path):
    dossier_reference = tmp_path / "reference"
    dossier_reference.mkdir()
    dossier_extractions = tmp_path / "extractions"
    dossier_extractions.mkdir()
    dossier_sortie = tmp_path / "sortie"

    _classeur_modele(dossier_reference)

    return {
        "reference": str(dossier_reference),
        "extractions": dossier_extractions,
        "sortie": str(dossier_sortie),
    }


def test_genere_un_nouveau_fichier_date_du_jour(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110", "37120"])

    resultat = generer_classeur(
        [str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29)
    )

    assert resultat["ok"] is True
    assert os.path.basename(resultat["chemin_genere"]) == "TRESORERIE JOURNALIÈRE et TDB DU  29 09 2026.xlsx"
    assert os.path.isfile(resultat["chemin_genere"])


def test_ne_modifie_jamais_le_modele_source(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110"])
    chemin_modele = os.path.join(contexte["reference"], "TRESORERIE JOURNALIÈRE et TDB DU  10 09 2026.xlsx")
    avant = open(chemin_modele, "rb").read()

    generer_classeur([str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    assert open(chemin_modele, "rb").read() == avant


def test_exclut_toujours_suivi_de_la_treso(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110"])

    resultat = generer_classeur([str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    assert "Suivi de la treso" not in classeur.sheetnames


def test_exclut_les_feuilles_mortes_de_lancienne_methode(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110"])

    resultat = generer_classeur([str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    for nom in ("Agences", "Dépôts", "Caisse", "SYNTHESE 2025", "Feuil1"):
        assert nom not in classeur.sheetnames
    assert classeur.sheetnames == ["Synthèse"]


def test_remplit_les_comptes_et_decale_le_j1(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110", "37120", "37220"])  # 2 Courants, 1 Chèques

    resultat = generer_classeur([str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    synthese = classeur["Synthèse"]
    assert synthese["C7"].value == 2  # Courants (37110+37120)
    assert synthese["C8"].value == 1  # Chèques (37220)
    assert synthese["C11"].value == 0  # Collectes : toujours 0
    assert synthese["C13"].value == 0  # Salariés : toujours 0
    assert synthese["C16"].value == "=SUM(C7:C15)"  # formule préservée, jamais remplacée
    assert synthese["C17"].value == 50  # J-1 = l'ancienne valeur de C16 (comptée hier)
    assert "Akwa" in resultat["agences_mises_a_jour"]


def test_agence_sans_fichier_reste_inchangee_mais_signalee(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110"])

    resultat = generer_classeur([str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    synthese = classeur["Synthèse"]
    assert synthese["F7"].value == 30  # Bafoussam inchangé (aucun fichier reçu)
    assert synthese["F17"].value == 30  # J-1 avancé quand même (ancienne valeur de F16, la formule)
    assert "Bafoussam" in resultat["agences_non_mises_a_jour"]
    assert "Bafoussam" not in resultat["agences_mises_a_jour"]


def test_ne_jamais_ecraser_un_fichier_deja_genere(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110"])

    premier = generer_classeur([str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))
    second = generer_classeur([str(chemin_akwa)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    assert premier["chemin_genere"] != second["chemin_genere"]
    assert os.path.isfile(premier["chemin_genere"])
    assert os.path.isfile(second["chemin_genere"])


def test_fichier_bloquant_est_ignore_pas_utilise(contexte):
    chemin_akwa = contexte["extractions"] / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_akwa, ["37110"])
    chemin_illisible = contexte["extractions"] / "Akwa_Compte_bis.xlsx"
    _extraction_comptes(chemin_illisible, ["37110"])  # même agence : le doublon devient bloquant

    resultat = generer_classeur(
        [str(chemin_akwa), str(chemin_illisible)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29)
    )

    assert "Akwa" not in resultat["agences_mises_a_jour"]
    assert set(resultat["fichiers_ignores"]) == {"Akwa_Compte.xlsx", "Akwa_Compte_bis.xlsx"}


def test_aucun_modele_disponible_est_signale_clairement(tmp_path):
    dossier_reference_vide = tmp_path / "reference_vide"
    dossier_reference_vide.mkdir()

    resultat = generer_classeur([], str(dossier_reference_vide), str(tmp_path / "sortie"))

    assert resultat["ok"] is False
    assert "modèle" in resultat["erreur"] or "référence" in resultat["erreur"]


def test_par_defaut_la_date_est_celle_dhier(contexte):
    """Corrige le bug signalé par l'utilisateur le 01/10/2026 : le classeur généré sans
    date explicite portait la date du jour d'exécution, alors que la trésorerie traitée
    chaque matin concerne la journée précédente."""
    resultat = generer_classeur([], contexte["reference"], contexte["sortie"])  # pas de `jour`

    attendu = date.today() - timedelta(days=1)
    assert resultat["date"] == attendu.isoformat()
    assert os.path.basename(resultat["chemin_genere"]) == (
        f"TRESORERIE JOURNALIÈRE et TDB DU  {attendu.day:02d} {attendu.month:02d} {attendu.year}.xlsx"
    )


def test_remplit_depots_et_engagements_et_decale_leur_j1(contexte):
    """Corrige le bug signalé par l'utilisateur le 01/10/2026 : les lignes 20/21 et 23/24
    du premier essai réel contenaient deux fois la même valeur (le N-1 n'était pas lu dans
    le classeur de référence)."""
    chemin_balance = contexte["extractions"] / "balance_akwa.pdf"
    _extraction_balance_classe3(chemin_balance, "DOUALA AKWA", depots=1_200_000, engagements=600_000)

    resultat = generer_classeur([str(chemin_balance)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    synthese = classeur["Synthèse"]
    assert synthese["C20"].value == 1_200_000  # nouveau dépôt Akwa
    assert synthese["C21"].value == 1_000_000  # J-1 = l'ancien dépôt du modèle (C20 avant écriture)
    assert synthese["C22"].value == "=C20-C21"  # formule préservée, jamais remplacée
    assert synthese["C23"].value == 600_000  # nouvel engagement Akwa
    assert synthese["C24"].value == 500_000  # J-1 = l'ancien engagement du modèle
    assert "Akwa" in resultat["agences_balance_mises_a_jour"]


def test_agence_sans_balance_garde_ses_valeurs_mais_j1_avance(contexte):
    chemin_balance = contexte["extractions"] / "balance_akwa.pdf"
    _extraction_balance_classe3(chemin_balance, "DOUALA AKWA", depots=1_200_000, engagements=600_000)

    resultat = generer_classeur([str(chemin_balance)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29))

    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    synthese = classeur["Synthèse"]
    # Bafoussam n'a reçu aucun fichier de balance : la valeur reste celle du modèle (2 000 000
    # / 800 000), mais le J-1 avance quand même vers cette même valeur (variation nulle, comme
    # pour les comptes).
    assert synthese["F20"].value == 2_000_000
    assert synthese["F21"].value == 2_000_000
    assert synthese["F23"].value == 800_000
    assert synthese["F24"].value == 800_000
    assert "Bafoussam" not in resultat["agences_balance_mises_a_jour"]


def test_nutilise_jamais_un_classeur_date_du_jour_genere_comme_son_propre_modele(tmp_path):
    """Corrige le bug du 01/10/2026 : si le dossier de référence contient déjà un classeur
    daté du jour qu'on génère (rattrapage, saisie en avance...), ce classeur ne doit jamais
    devenir son propre modèle (sinon son propre J-1 deviendrait une référence de lui-même)."""
    dossier_reference = tmp_path / "reference"
    dossier_reference.mkdir()
    _classeur_modele(dossier_reference, "TRESORERIE JOURNALIÈRE et TDB DU  29 09 2026.xlsx")
    # Un classeur daté du jour qu'on s'apprête à générer existe déjà (ex. rempli à la main) :
    classeur_meme_jour = openpyxl.Workbook()
    classeur_meme_jour.active.title = "Synthèse"
    classeur_meme_jour["Synthèse"]["C20"] = 999_999_999  # ne doit jamais se retrouver en J-1
    classeur_meme_jour.save(dossier_reference / "TRESORERIE JOURNALIÈRE et TDB DU  30 09 2026.xlsx")

    resultat = generer_classeur([], str(dossier_reference), str(tmp_path / "sortie"), jour=date(2026, 9, 30))

    assert resultat["ok"] is True
    assert "29 09 2026" in resultat["modele_utilise"]
    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    assert classeur["Synthèse"]["C21"].value != 999_999_999


def test_fichier_balance_bloquant_nest_jamais_utilise(contexte):
    chemin_1 = contexte["extractions"] / "balance_akwa_1.pdf"
    _extraction_balance_classe3(chemin_1, "DOUALA AKWA", depots=1_200_000, engagements=600_000)
    chemin_2 = contexte["extractions"] / "balance_akwa_2.pdf"
    _extraction_balance_classe3(chemin_2, "DOUALA AKWA", depots=1_300_000, engagements=650_000)

    resultat = generer_classeur(
        [str(chemin_1), str(chemin_2)], contexte["reference"], contexte["sortie"], jour=date(2026, 9, 29)
    )

    classeur = openpyxl.load_workbook(resultat["chemin_genere"])
    synthese = classeur["Synthèse"]
    assert synthese["C20"].value == 1_000_000  # inchangé : les deux fichiers en doublon sont ignorés
    assert "Akwa" not in resultat["agences_balance_mises_a_jour"]
    assert set(resultat["fichiers_ignores"]) == {"balance_akwa_1.pdf", "balance_akwa_2.pdf"}
