"""Tests de l'export du rapport d'analyse en classeur Excel (10/10/2026)."""

import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from orisflow_engine.rapport import exporter_rapport_excel


def _fichier(**kwargs):
    base = {
        "nom": "Akwa_Compte.xlsx",
        "type_libelle": "Liste de comptes",
        "agence_libelle": "Akwa",
        "confiance_agence": "nom",
        "total_comptes": 123,
        "niveau": "information",
        "messages": [],
    }
    base.update(kwargs)
    return base


def test_exporte_une_ligne_par_fichier(tmp_path):
    chemin = tmp_path / "Rapport-analyse.xlsx"
    fichiers = [
        _fichier(),
        _fichier(nom="Mokolo_Compte.xlsx", agence_libelle="Mokolo", niveau="bloquant", messages=["Doublon."]),
    ]

    exporter_rapport_excel(fichiers, str(chemin))

    assert chemin.is_file()
    classeur = openpyxl.load_workbook(chemin)
    feuille = classeur["Rapport d'analyse"]
    entetes = [feuille.cell(row=1, column=c).value for c in range(1, 8)]
    assert entetes == ["Fichier", "Type détecté", "Agence", "Confiance agence", "Comptes", "Niveau", "Messages"]
    assert feuille.cell(row=2, column=1).value == "Akwa_Compte.xlsx"
    assert feuille.cell(row=2, column=6).value == "Conforme"
    assert feuille.cell(row=3, column=1).value == "Mokolo_Compte.xlsx"
    assert feuille.cell(row=3, column=6).value == "Rejeté"
    assert feuille.cell(row=3, column=7).value == "Doublon."


def test_journal_etapes_precede_le_tableau(tmp_path):
    chemin = tmp_path / "Rapport-analyse.xlsx"
    exporter_rapport_excel([_fichier()], str(chemin), journal_etapes=["1 agence confirmée manuellement."])

    classeur = openpyxl.load_workbook(chemin)
    feuille = classeur["Rapport d'analyse"]
    assert feuille.cell(row=1, column=1).value == "Étapes de l'analyse"
    assert feuille.cell(row=2, column=1).value == "- 1 agence confirmée manuellement."
    # Ligne vide puis en-tête du tableau.
    assert feuille.cell(row=4, column=1).value == "Fichier"


def test_naccrase_pas_mais_le_dossier_parent_est_cree(tmp_path):
    chemin = tmp_path / "sous_dossier" / "Rapport.xlsx"
    exporter_rapport_excel([], str(chemin))
    assert chemin.is_file()
