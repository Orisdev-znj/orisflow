"""Module « Téléverser sur CloudBank » (10/10/2026) : mapping, résolution par agence, écriture, import.

Aucune donnée réelle : les fichiers d'essai sont fabriqués ici, les référentiels sont ceux
embarqués par défaut (comptes du plan, règles de mapping)."""

import io
import json
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from orisflow_engine import cli
from orisflow_engine import cloudbank_extourne
from orisflow_engine.cloudbank_ecriture import (
    COMPTE_CHARGES_DIVERSES, chemin_sans_ecrasement, consolider_charges_diverses, ecrire_fichier)
from orisflow_engine.cloudbank_extraire import extraire
from orisflow_engine.cloudbank_import import deposer_fichier, lire_fichier_mapping, lire_nom_mapping
from orisflow_engine.cloudbank_mapping import mapper_ligne, rechercher
from orisflow_engine.cloudbank_referentiels import (
    ErreurCloudBank, Referentiels, ajouter_mapping_valide, dossier_donnees_par_defaut, preparer_references)
from orisflow_engine.cloudbank_resolution_agence import resoudre_compte_agence


@pytest.fixture(scope="module")
def ref():
    return Referentiels()


def executer(commande, parametres, monkeypatch, capsys):
    monkeypatch.setattr(sys, "stdin", io.StringIO(json.dumps(parametres)))
    code = cli.main([commande])
    lignes = [json.loads(l) for l in capsys.readouterr().out.splitlines() if l.strip()]
    return code, lignes


# --- Référentiels -------------------------------------------------------------------------------

def test_les_donnees_par_defaut_sont_embarquees():
    assert os.path.isfile(os.path.join(dossier_donnees_par_defaut(), "mappings_valides.json"))


def test_la_copie_modifiable_est_creee_sans_ecraser_une_copie_existante(tmp_path):
    dossier = str(tmp_path / "CloudBank" / "Reférences")
    preparer_references(dossier)
    modifiee = os.path.join(dossier, "agences.json")
    with open(modifiee, "w", encoding="utf-8") as f:
        f.write('{"10001": "AKWA MODIFIEE"}')
    preparer_references(dossier)
    assert "AKWA MODIFIEE" in open(modifiee, encoding="utf-8").read()


def test_ajouter_un_mapping_ne_touche_que_la_copie_modifiable(tmp_path):
    defaut = os.path.join(dossier_donnees_par_defaut(), "mappings_valides.json")
    avant = open(defaut, encoding="utf-8").read()
    ref = Referentiels(str(tmp_path / "Ref"))
    assert ajouter_mapping_valide(ref, "Achat de bougies d'essai", "6523600000201", "Entretien")
    assert not ajouter_mapping_valide(ref, "Achat de bougies d'essai", "6523600000201", "Entretien")
    assert open(defaut, encoding="utf-8").read() == avant
    assert Referentiels(str(tmp_path / "Ref")).mappings


def test_sans_copie_modifiable_aucun_enregistrement(ref):
    assert not ajouter_mapping_valide(ref, "Libellé inventé", "6523600000201", "x")


# --- Mapping ------------------------------------------------------------------------------------

def test_un_libelle_valide_est_reconnu_tel_quel(ref):
    resultat = mapper_ligne(ref, "Frais de transport")
    assert resultat["statut"] == "valide" and resultat["compte"] == "6522000000001"


def test_un_libelle_inconnu_reste_non_trouve_avec_suggestions(ref):
    resultat = mapper_ligne(ref, "zzzz qqqq xxxx")
    assert resultat["statut"] == "non_trouve" and resultat["compte"] is None


def test_provision_n_est_pas_reconnue_dans_approvisionnement(ref):
    """Faux positif réel du skill (Kousseri) : un mot-clé ne doit pas matcher au milieu d'un mot."""
    resultat = mapper_ligne(ref, "approvisionnement caisse")
    assert resultat["statut"] != "valide"


def test_la_recherche_trouve_des_comptes(ref):
    resultat = rechercher(ref, "transport")
    assert resultat["cloudbank"] or resultat["pcemf"]


# --- Résolution du compte par agence (règle n°7) -------------------------------------------------

def test_le_siege_garde_le_compte_generique(ref):
    famille = next(iter(ref.comptes_par_agence))
    assert resoudre_compte_agence(ref, famille, 10000) == (famille, None)


def test_un_compte_confirme_pour_l_agence_est_remplace_sans_avertissement(ref):
    famille, variantes = next((f, v) for f, v in ref.comptes_par_agence.items() if v)
    agence, reel = next(iter(variantes.items()))
    assert resoudre_compte_agence(ref, famille, agence) == (reel, None)


def test_un_compte_gere_non_confirme_pour_l_agence_avertit(ref):
    """Un compte confirmé pour d'autres agences n'est pas confirmé pour celle-ci : avertissement,
    jamais de compte deviné (le 468 de Balessing et de Kousseri valent le même numéro par coïncidence)."""
    famille, agence = next(
        (f, a) for f in sorted(ref.comptes_geres) for a in ref.agences
        if a != "10000" and f not in ref.comptes_deja_resolus_par_agence.get(a, ()) and a not in ref.comptes_par_agence.get(f, {}))
    compte, avertissement = resoudre_compte_agence(ref, famille, int(agence))
    assert compte == famille and avertissement and "non confirmé" in avertissement


def test_un_compte_externe_complet_ne_declenche_jamais_d_avertissement(ref):
    assert resoudre_compte_agence(ref, "3731234567890", 30001) == ("3731234567890", None)


# --- Écriture -----------------------------------------------------------------------------------

def _charges():
    return [
        {"compte": "6522000000001", "intitule": "Transports", "libelle": "Taxi", "montant": 3500},
        {"compte": "6522000000001", "intitule": "Transports", "libelle": "Bus", "montant": 1500},
        {"compte": COMPTE_CHARGES_DIVERSES, "intitule": "Charges diverses", "libelle": "Divers A", "montant": 1000},
        {"compte": COMPTE_CHARGES_DIVERSES, "intitule": "Charges diverses", "libelle": "Divers B", "montant": 2000},
    ]


def test_petite_caisse_equilibree_sur_les_deux_feuilles(ref, tmp_path):
    rapport = ecrire_fichier(ref, "petite_caisse", _charges(), str(tmp_path), agence=10000, mois="Septembre", annee="2026")
    assert [f["feuille"] for f in rapport["feuilles"]] == ["ECRITURE PAIE DG", "SYNTHESE PAR COMPTE"]
    assert all(f["ecart"] == 0 and f["total_debit"] == 8000 for f in rapport["feuilles"])
    assert os.path.basename(rapport["fichier"]) == "ECRITURE PETITE CAISSE A TELEVERSER - SIEGE - Septembre 2026.xlsx"


def test_charges_diverses_consolidees_dans_le_detail(ref):
    lignes = [{"compte": l["compte"], "libelle": l["libelle"], "debit": l["montant"]} for l in _charges()]
    consolidees = consolider_charges_diverses(lignes)
    diverses = [l for l in consolidees if l["compte"] == COMPTE_CHARGES_DIVERSES]
    assert len(diverses) == 1 and diverses[0]["debit"] == 3000 and "Divers A" in diverses[0]["libelle"]


def test_retour_de_caisse_ajoute_deux_lignes_equilibrees(ref, tmp_path):
    rapport = ecrire_fichier(ref, "petite_caisse", _charges(), str(tmp_path), agence=10000,
                             mois="Septembre", annee="2026", retour=500)
    assert all(f["ecart"] == 0 for f in rapport["feuilles"])
    assert rapport["feuilles"][0]["lignes"] == 3 + 2 + 1  # 3 charges (diverses fusionnées) + retour + contrepartie


def test_un_fichier_existant_n_est_jamais_ecrase(ref, tmp_path):
    a = ecrire_fichier(ref, "petite_caisse", _charges(), str(tmp_path), mois="Septembre", annee="2026")["fichier"]
    b = ecrire_fichier(ref, "petite_caisse", _charges(), str(tmp_path), mois="Septembre", annee="2026")["fichier"]
    assert a != b and os.path.exists(a) and os.path.exists(b) and b.endswith("(2).xlsx")


def test_une_ligne_sans_compte_empeche_la_generation(ref, tmp_path):
    with pytest.raises(ErreurCloudBank, match="pas de compte"):
        ecrire_fichier(ref, "petite_caisse", [{"compte": None, "montant": 10}], str(tmp_path), mois="Septembre", annee="2026")


def test_un_compte_mal_forme_empeche_la_generation(ref, tmp_path):
    with pytest.raises(ErreurCloudBank, match="13 chiffres"):
        ecrire_fichier(ref, "salaires", [{"compte": "123", "debit": 10, "agence": 10000}], str(tmp_path))


def test_le_dossier_de_sortie_trop_long_est_accepte(ref, tmp_path):
    long = tmp_path
    for _ in range(6):
        long = long / ("x" * 45)
    rapport = ecrire_fichier(ref, "petite_caisse", _charges(), str(long), mois="Septembre", annee="2026")
    assert len(rapport["fichier"]) > 260
    from orisflow_engine.chemins import chemin_lecture
    assert os.path.exists(chemin_lecture(rapport["fichier"]))


# --- Import du fichier « MAPPING A VALIDER » ----------------------------------------------------

def _mapping(chemin, lignes, entetes=("LIBELLE", "MONTANT", "COMPTE PROPOSE", "INTITULE", "STATUT", "SOURCE", "COMMENTAIRE")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(entetes))
    for ligne in lignes:
        ws.append(list(ligne))
    wb.save(chemin)


def test_import_lit_les_lignes_et_l_agence_du_nom(ref, tmp_path):
    chemin = str(tmp_path / "MAPPING A VALIDER - MARCHE CENTRAL - Septembre 2026.xlsx")
    _mapping(chemin, [
        ("Frais de transport", 3500, "6522000000001", "Transports", "valide", "écriture validée", None),
        ("Achat registre", 2000, "6521200000119", "Matériel", "à confirmer", "règle mot-clé", None),
        ("Achat un ballot SITA", 11500, None, None, "non_trouve", None, None),
    ])
    donnees = lire_fichier_mapping(ref, chemin)
    assert donnees["contexte"]["agence"] == "20002" and donnees["contexte"]["mois"] == "Septembre"
    assert donnees["resume"] == {"valide": 1, "a_confirmer": 1, "non_trouve": 1}
    assert donnees["total"] == 17000
    assert donnees["lignes"][2]["compte"] is None and donnees["lignes"][2]["suggestions"] is not None


def test_import_refuse_un_fichier_sans_les_colonnes_attendues(ref, tmp_path):
    chemin = str(tmp_path / "x.xlsx")
    _mapping(chemin, [("a", 1)], entetes=("LIBELLE", "MONTANT"))
    with pytest.raises(ErreurCloudBank, match="COMPTE PROPOSE"):
        lire_fichier_mapping(ref, chemin)


def test_import_ne_modifie_pas_le_fichier_source(ref, tmp_path):
    chemin = str(tmp_path / "MAPPING A VALIDER - AKWA - Septembre 2026.xlsx")
    _mapping(chemin, [("Frais de transport", 100, "6522000000001", "T", "valide", "", None)])
    avant = open(chemin, "rb").read()
    lire_fichier_mapping(ref, chemin)
    assert open(chemin, "rb").read() == avant


def test_un_nom_hors_convention_est_signale_sans_bloquer(ref):
    assert lire_nom_mapping(ref, "mon fichier.xlsx")["conforme"] is False


def test_un_compte_non_conforme_devient_non_trouve(ref, tmp_path):
    chemin = str(tmp_path / "MAPPING A VALIDER - AKWA - Septembre 2026.xlsx")
    _mapping(chemin, [("Achat X", 100, "12345", "T", "valide", "", None)])
    ligne = lire_fichier_mapping(ref, chemin)["lignes"][0]
    assert ligne["statut"] == "non_trouve" and ligne["compte"] is None and "13 chiffres" in ligne["avertissement"]


# --- Extourne, extraction, dépôt ----------------------------------------------------------------

def test_extourne_un_debit_par_transaction_et_un_credit_cumule(ref, tmp_path, monkeypatch):
    monkeypatch.setattr(cloudbank_extourne, "lire_historique_compte", lambda chemin: [
        {"libelle": "FRAIS TENUE", "debit": 500.0, "credit": 0.0},
        {"libelle": "FRAIS TENUE", "debit": 500.0, "credit": 0.0},
        {"libelle": "AUTRE", "debit": 20.0, "credit": 0.0},
    ])
    rapport = cloudbank_extourne.extourner(
        ref, "x.xls", [{"montant": 500, "compte": "7200000000001", "intitule": "Frais"}],
        "3731234567890", "CLIENT TEST", 30001, str(tmp_path))
    assert rapport["transactions_extournees"] == 2 and rapport["cumul_credite"] == 1000 and rapport["ecart"] == 0
    feuille = openpyxl.load_workbook(rapport["fichier"]).active
    assert feuille.cell(2, 3).value.startswith("EXT ") and not str(feuille.cell(4, 3).value).startswith("EXT ")


def test_extourne_sans_transaction_correspondante_est_une_erreur_claire(ref, tmp_path, monkeypatch):
    monkeypatch.setattr(cloudbank_extourne, "lire_historique_compte", lambda chemin: [{"libelle": "A", "debit": 1.0, "credit": 0.0}])
    with pytest.raises(ErreurCloudBank, match="aucune transaction"):
        cloudbank_extourne.extourner(ref, "x.xls", [{"montant": 500, "compte": "7200000000001"}],
                                     "3731234567890", "C", 30001, str(tmp_path))


def _grand_livre(chemin):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(6):
        ws.append(["", "en-tête général"])
    ws.append(["", "Agence  :  30001---BALESSING"])
    ws.append(["", "65220 : TRANSPORTS"])
    ws.append(["", "Agence", "Date", "Réf", "Libellé"])
    ws.append(["", "6522000000009 : TRANSPORT COMMERCIAUX"])
    ws.append(["", "30001", "01/09", "R1", "Taxi", 1500])
    ws.append(["", "", "", "", "", "TOTAL COMPTE", "", 1500])
    ws.append(["", "6522000000010 : AUTRE COMPTE"])
    ws.append(["", "30001", "02/09", "R2", "Bus", 900])
    ws.append(["", "", "", "", "", "TOTAL COMPTE", "", 900])
    wb.save(chemin)


def test_extraire_ne_garde_que_les_sous_comptes_demandes(tmp_path):
    source = str(tmp_path / "GL.xlsx")
    _grand_livre(source)
    avant = open(source, "rb").read()
    rapport = extraire(source, ["transport commerciaux"], str(tmp_path / "sortie"))
    assert rapport["nombre_sous_comptes"] == 1 and rapport["somme_total_compte"] == 1500
    texte = " ".join(str(c.value) for ligne in openpyxl.load_workbook(rapport["fichier"]).active.iter_rows() for c in ligne)
    assert "TRANSPORT COMMERCIAUX" in texte and "AUTRE COMPTE" not in texte
    assert open(source, "rb").read() == avant


def test_extraire_sans_terme_est_une_erreur_claire(tmp_path):
    with pytest.raises(ErreurCloudBank, match="au moins un terme"):
        extraire("x.xlsx", ["  "], str(tmp_path))


def test_deposer_ne_remplace_jamais_un_fichier_existant(tmp_path):
    source = str(tmp_path / "EXT A TELEVERSER.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["NUMERO DE COMPTE", "INTITULE", "LIBELLE", "DEBIT ", "CREDIT", "PROVISION ", "AGENCE"])
    ws.append(["6522000000001", "T", "EXT x", 500, None, 1, 30001])
    ws.append(["3731234567890", "C", "x", None, 500, 0, 30001])
    wb.save(source)
    a = deposer_fichier(source, str(tmp_path / "Resultats"))
    b = deposer_fichier(source, str(tmp_path / "Resultats"))
    assert a["fichier"] != b["fichier"] and a["feuilles"][0]["ecart"] == 0 and a["feuilles"][0]["lignes_ecriture"] == 2


# --- Protocole CLI ------------------------------------------------------------------------------

def test_cli_ecrire_puis_importer(monkeypatch, capsys, tmp_path):
    code, messages = executer("cloudbank_ecrire", {
        "modele": "petite_caisse", "lignes": _charges(), "agence": 10000, "mois": "Septembre", "annee": "2026",
        "dossierSortie": str(tmp_path), "dossierReferences": str(tmp_path / "Ref")}, monkeypatch, capsys)
    assert code == 0 and messages[-1]["type"] == "resultat" and os.path.exists(messages[-1]["fichier"])
    assert os.path.isdir(tmp_path / "Ref")


def test_cli_renvoie_une_erreur_lisible_sans_trace_python(monkeypatch, capsys, tmp_path):
    _, messages = executer("cloudbank_ecrire", {
        "modele": "petite_caisse", "lignes": [{"compte": None, "montant": 5}], "mois": "Septembre", "annee": "2026",
        "dossierSortie": str(tmp_path)}, monkeypatch, capsys)
    assert messages[-1]["type"] == "erreur" and "Traceback" not in messages[-1]["message"]


def test_cli_confirmer_ligne_refuse_un_numero_a_saisir_a_la_main(monkeypatch, capsys):
    _, messages = executer("cloudbank_confirmer_ligne", {"compte": "12"}, monkeypatch, capsys)
    assert messages[-1]["type"] == "erreur" and "13 chiffres" in messages[-1]["message"]


def test_cli_confirmer_ligne_n_enregistre_que_sur_demande(monkeypatch, capsys, tmp_path):
    parametres = {"compte": "6522000000001", "libelle": "Un libellé tout neuf", "dossierReferences": str(tmp_path / "Ref")}
    _, sans = executer("cloudbank_confirmer_ligne", parametres, monkeypatch, capsys)
    assert sans[-1]["enregistre"] is False
    _, avec = executer("cloudbank_confirmer_ligne", {**parametres, "enregistrer": True}, monkeypatch, capsys)
    assert avec[-1]["enregistre"] is True


def test_cli_importer_un_fichier_introuvable(monkeypatch, capsys):
    _, messages = executer("cloudbank_importer_mapping", {"chemin": "absent.xlsx"}, monkeypatch, capsys)
    assert messages[-1]["type"] == "erreur"
