import os
import sys

import openpyxl
import pymupdf

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from orisflow_engine.classification import classer_fichiers


def _extraction_comptes(chemin, prefixes):
    """Fabrique une extraction de comptes avec des numéros au format valide
    (5 chiffres-6 chiffres-2 chiffres), à partir d'une liste de préfixes à 5 chiffres."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    for _ in range(23):
        feuille.append([])
    feuille.append(["N°", "Numero de compte"])
    for i, prefixe in enumerate(prefixes, start=1):
        feuille.append([i, f"{prefixe}-{i:06d}-00"])
    classeur.save(chemin)


def _classeur_reference(dossier, valeurs_ligne16):
    """Fabrique un classeur minimal « TRESORERIE JOURNALIÈRE... » avec une feuille Synthèse."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "Synthèse"
    for colonne, valeur in valeurs_ligne16.items():
        feuille[f"{colonne}16"] = valeur
    chemin = dossier / "TRESORERIE JOURNALIÈRE et TDB DU  10 09 2026.xlsx"
    classeur.save(chemin)
    return dossier


def test_reconnait_type_et_agence_dun_fichier_de_comptes(tmp_path):
    chemin = tmp_path / "Bafoussam_Compte.xlsx"
    _extraction_comptes(chemin, ["37110", "37120"])

    resultat = classer_fichiers([str(chemin)])

    assert resultat["ok"] is True
    fichier = resultat["fichiers"][0]
    assert fichier["type_detecte"] == "compte"
    assert fichier["agence_detectee"] == "bafoussam"
    assert fichier["total_comptes"] == 2
    assert fichier["niveau"] == "information"


def test_fichier_introuvable_est_bloquant():
    resultat = classer_fichiers([r"C:\ceci\nexiste\pas.xlsx"])
    assert resultat["ok"] is False
    assert resultat["fichiers"][0]["niveau"] == "bloquant"


def test_agence_inconnue_est_un_avertissement(tmp_path):
    chemin = tmp_path / "export_du_jour.xlsx"
    _extraction_comptes(chemin, ["37110"])
    resultat = classer_fichiers([str(chemin)])
    fichier = resultat["fichiers"][0]
    assert fichier["agence_detectee"] is None
    assert fichier["niveau"] == "avertissement"


def test_deux_fichiers_pour_la_meme_agence_sont_bloquants(tmp_path):
    """Contenus réellement différents (pas un simple doublon) : Orisflow ne peut pas deviner
    lequel des deux est le bon pour cette agence, le blocage reste nécessaire pour les deux.
    (Avant le 10/10/2026, ce test utilisait deux fichiers au contenu strictement identique :
    ce cas est désormais couvert séparément, et n'est plus bloquant — voir
    `test_doublon_au_contenu_identique_sur_meme_agence_ne_bloque_pas_le_fichier_conserve`.)
    37110 et 37120 comptent tous les deux comme « courants » : il faut une autre catégorie
    (37420, garanties) pour obtenir des comptages réellement différents."""
    chemin1 = tmp_path / "Akwa_Compte.xlsx"
    chemin2 = tmp_path / "Akwa_Compte_bis.xlsx"
    _extraction_comptes(chemin1, ["37110"])
    _extraction_comptes(chemin2, ["37420"])

    resultat = classer_fichiers([str(chemin1), str(chemin2)])

    assert resultat["ok"] is False
    assert all(f["niveau"] == "bloquant" for f in resultat["fichiers"])


def test_ecart_anormal_avec_la_veille_suggere_lagence_probable(tmp_path):
    dossier_extractions = tmp_path / "extractions"
    dossier_extractions.mkdir()
    dossier_reference = tmp_path / "reference"
    dossier_reference.mkdir()

    # Le fichier nommé « Bafoussam » contient en réalité le total de Balessing (cas du 10/09/2026).
    chemin = dossier_extractions / "Bafoussam_Compte.xlsx"
    _extraction_comptes(chemin, ["37420"] * 5)  # 5 comptes Garanties
    _classeur_reference(dossier_reference, {"F": 2441, "H": 5})  # F=Bafoussam, H=Balessing

    resultat = classer_fichiers([str(chemin)], dossier_reference=str(dossier_reference))

    fichier = resultat["fichiers"][0]
    assert fichier["niveau"] == "avertissement"
    assert any("Balessing" in m for m in fichier["messages"])
    assert resultat["reference"]["disponible"] is True


def test_numero_repete_dans_une_liste_nest_pas_un_avertissement(tmp_path):
    chemin = tmp_path / "Akwa_Compte.xlsx"
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    for _ in range(23):
        feuille.append([])
    feuille.append(["N°", "Numero de compte"])
    feuille.append([1, "37110-000001-00"])
    feuille.append([2, "37110-000001-00"])  # même numéro que la ligne précédente
    classeur.save(chemin)

    resultat = classer_fichiers([str(chemin)])
    fichier = resultat["fichiers"][0]
    assert fichier["doublons"] == ["37110-000001-00"]
    # Décision du 03/10/2026 : les numéros répétés dans une liste ne sont plus un avertissement.
    assert fichier["niveau"] == "information"
    assert not any("double" in m for m in fichier["messages"])  # plus de message


def test_signale_les_numeros_mal_formes(tmp_path):
    chemin = tmp_path / "Akwa_Compte.xlsx"
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    for _ in range(23):
        feuille.append([])
    feuille.append(["N°", "Numero de compte"])
    feuille.append([1, "37110-000001-00"])
    feuille.append([2, "371100000100"])  # sans les tirets attendus

    classeur.save(chemin)

    resultat = classer_fichiers([str(chemin)])
    fichier = resultat["fichiers"][0]
    assert fichier["mal_formes"] == ["371100000100"]
    assert fichier["niveau"] == "avertissement"


def test_reconnait_une_balance_classe3_et_lit_agence_depuis_le_contenu(tmp_path):
    chemin = tmp_path / "EtBalance_General_Consolide_0000847.pdf"  # nom sans agence, comme les vrais fichiers
    document = pymupdf.open()
    page = document.new_page()
    ancres = [231, 291, 351, 411, 471, 531]
    page.insert_text((10, 50), "Balance generale consolidée   Chapitre de : 3  à : 3")
    page.insert_text((10, 70), "Groupe: DOUALA AKWA   DOUALA AKWA")
    page.insert_text((10, 150), "Compte")
    page.insert_text((71, 150), "Intitulé")
    for i, x in enumerate(ancres):
        page.insert_text((x, 150), "Debit" if i % 2 == 0 else "Crédit")
    page.insert_text((10, 175), "Total Classe : 3 TOTAL")
    for x, valeur in zip(ancres, [1, 2, 3, 4, 500, 600]):
        page.insert_text((x + 21, 175), str(valeur))
    document.save(str(chemin))
    document.close()

    resultat = classer_fichiers([str(chemin)])
    fichier = resultat["fichiers"][0]

    assert fichier["type_detecte"] == "balance_classe3"
    assert fichier["agence_detectee"] == "akwa"  # trouvée dans le contenu, pas dans le nom du fichier
    assert fichier["depots"] == 600
    assert fichier["engagements"] == 500


def test_reconnait_un_releve_bancaire_pdf(tmp_path):
    chemin = tmp_path / "releve.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (50, 50),
        "Code client : 735378\nEXTRAIT DE COMPTE\nNumero de compte : 10038-01773537801-12\n"
        "Solde (XAF) au 02/10/2026 : 186 412 276",
    )
    document.save(str(chemin))
    document.close()

    resultat = classer_fichiers([str(chemin)])

    fichier = resultat["fichiers"][0]
    assert fichier["type_detecte"] == "releve_cca"
    assert fichier["numero_compte_pdf"] == "10038-01773537801-12"
    # Depuis le 02/10/2026 : la clé RIB (12) identifie l'agence (Akwa), voir regles_banques.py.
    assert fichier["agence_detectee"] == "akwa"
    assert fichier["ligne_banque_cible"] == "cca_bank"
    assert fichier["solde_releve"] == 186412276


# --- Reconnaissance par comptage (démarré le 03/10/2026) -----------------------------


def test_agence_deduite_par_comptage_en_dernier_recours(tmp_path):
    """Ni le nom ni les numéros de compte ne donnent l'agence : Orisflow compare le total
    de comptes à celui de la veille (démarré le 03/10/2026)."""
    dossier_reference = tmp_path / "reference"
    dossier_reference.mkdir()
    _classeur_reference(dossier_reference, {"C": 1})  # Akwa (colonne C), total veille = 1

    chemin = tmp_path / "ETListeCompte_NoHeader_0006870.xlsx"
    _extraction_comptes(chemin, ["37110"])  # 1 seul compte, sans agence dans le nom

    resultat = classer_fichiers([str(chemin)], dossier_reference=str(dossier_reference))

    fichier = resultat["fichiers"][0]
    assert fichier["agence_detectee"] == "akwa"
    assert fichier["confiance_agence"] == "comptage"
    assert fichier["niveau"] == "avertissement"  # à vérifier, jamais une certitude
    assert "proximité" in fichier["messages"][-1].lower()


def test_comptage_evite_les_conflits_quand_on_importe_les_12_listes_ensemble(tmp_path):
    """L'utilisateur importe toujours les 12 listes en même temps (demande du 03/10/2026) :
    deux fichiers sans agence reconnue ne doivent jamais se voir attribuer la même agence."""
    dossier_reference = tmp_path / "reference"
    dossier_reference.mkdir()
    _classeur_reference(dossier_reference, {"C": 100, "D": 50})  # Akwa=100, Mokolo=50

    chemin_1 = tmp_path / "ETListeCompte_NoHeader_0006870.xlsx"
    _extraction_comptes(chemin_1, ["37110"] * 101)  # proche d'Akwa (100) ET de rien d'autre
    chemin_2 = tmp_path / "ETListeCompte_NoHeader_0006871.xlsx"
    _extraction_comptes(chemin_2, ["37110"] * 49)  # proche de Mokolo (50)

    resultat = classer_fichiers([str(chemin_1), str(chemin_2)], dossier_reference=str(dossier_reference))

    agences = {f["nom"]: f["agence_detectee"] for f in resultat["fichiers"]}
    assert agences["ETListeCompte_NoHeader_0006870.xlsx"] == "akwa"
    assert agences["ETListeCompte_NoHeader_0006871.xlsx"] == "mokolo"


def test_confirmation_manuelle_est_prioritaire_sur_tout(tmp_path):
    chemin = tmp_path / "ETListeCompte_NoHeader_0006870.xlsx"
    _extraction_comptes(chemin, ["37110"])

    resultat = classer_fichiers(
        [str(chemin)], agences_manuelles={str(chemin): "bepanda"},
    )

    fichier = resultat["fichiers"][0]
    assert fichier["agence_detectee"] == "bepanda"
    assert fichier["confiance_agence"] == "manuelle"
    assert fichier["niveau"] == "information"


def test_journal_et_progression_par_fichier(tmp_path):
    """Chaque fichier classé déclenche le callback (journal visible pendant l'analyse)."""
    chemin_1 = tmp_path / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin_1, ["37110"])
    chemin_2 = tmp_path / "Export_Compte_du_jour.xlsx"
    _extraction_comptes(chemin_2, ["37120"])  # contenu différent : pas un doublon

    vus = []
    resultat = classer_fichiers([str(chemin_1), str(chemin_2)], sur_fichier_classe=vus.append)

    assert [f["nom"] for f in vus] == ["Akwa_Compte.xlsx", "Export_Compte_du_jour.xlsx"]
    assert resultat["journal_etapes"] == [
        "1 fichier(s) encore sans agence : votre confirmation sera demandée."
    ]


def test_fichiers_au_contenu_identique_sont_des_doublons(tmp_path):
    """Décision du 03/10/2026 : doublon = contenu strictement identique, quel que soit le nom."""
    original = tmp_path / "Akwa_Compte.xlsx"
    _extraction_comptes(original, ["37110"])
    copie = tmp_path / "Export_Compte_copie.xlsx"
    copie.write_bytes(original.read_bytes())

    resultat = classer_fichiers([str(original), str(copie)])

    premier, second = resultat["fichiers"]
    assert premier["niveau"] != "bloquant"
    assert second["niveau"] == "bloquant"
    assert "identique" in second["messages"][-1]
    assert resultat["ok"] is False


def test_fichiers_differents_ne_sont_pas_des_doublons(tmp_path):
    a = tmp_path / "Akwa_Compte.xlsx"
    _extraction_comptes(a, ["37110"])
    b = tmp_path / "Mokolo_Compte.xlsx"
    _extraction_comptes(b, ["37120"])

    resultat = classer_fichiers([str(a), str(b)])

    assert all(f["niveau"] != "bloquant" for f in resultat["fichiers"])


def test_doublon_au_contenu_identique_sur_meme_agence_ne_bloque_pas_le_fichier_conserve(tmp_path):
    """Trouvé le 10/10/2026 sur un vrai lot d'export (Ndogpassi exporté 4 fois par erreur) :
    même quand l'agence est reconnue avec certitude pour plusieurs fichiers identiques,
    `detecter_doublons` ré-escaladait aussi le fichier conservé par `detecter_fichiers_identiques`,
    bloquant l'agence entière alors qu'un seul exemplaire aurait suffi. Ici, deux fichiers au
    contenu strictement identique forcés sur la même agence (confirmation manuelle, comme le
    ferait une identification par numéros de compte à 15/15) : seul le second doit rester
    bloquant, le premier doit rester utilisable."""
    original = tmp_path / "ETListeCompte_NoHeader_0001.xlsx"
    _extraction_comptes(original, ["37110"])
    copie = tmp_path / "ETListeCompte_NoHeader_0002.xlsx"
    copie.write_bytes(original.read_bytes())

    resultat = classer_fichiers(
        [str(original), str(copie)],
        agences_manuelles={str(original): "ndogpassi", str(copie): "ndogpassi"},
    )

    premier, second = resultat["fichiers"]
    assert premier["agence_detectee"] == "ndogpassi"
    assert premier["niveau"] != "bloquant"
    assert "Plusieurs fichiers correspondent" not in " ".join(premier["messages"])
    assert second["niveau"] == "bloquant"
    assert "identique" in " ".join(second["messages"])


def test_doublon_au_contenu_different_sur_meme_agence_reste_bloquant(tmp_path):
    """Deux fichiers réellement différents (pas de simple doublon de contenu, ni de valeurs
    identiques) mais résolus sur la même agence : Orisflow ne peut pas deviner lequel est le
    bon, le blocage reste nécessaire pour les deux. 37110 et 37120 comptent tous les deux comme
    « courants » : il faut 37420 (garanties) pour obtenir des comptages réellement différents."""
    a = tmp_path / "ETListeCompte_NoHeader_0003.xlsx"
    _extraction_comptes(a, ["37110"])
    b = tmp_path / "ETListeCompte_NoHeader_0004.xlsx"
    _extraction_comptes(b, ["37420"])

    resultat = classer_fichiers(
        [str(a), str(b)],
        agences_manuelles={str(a): "ndogpassi", str(b): "ndogpassi"},
    )

    assert all(f["niveau"] == "bloquant" for f in resultat["fichiers"])
    assert all("valeurs différentes" in " ".join(f["messages"]) for f in resultat["fichiers"])


def test_releves_manquants_proposent_la_valeur_de_la_veille(tmp_path):
    """Aucun relevé reçu aujourd'hui : chaque relevé attendu est listé, avec la valeur du carnet
    de la veille quand elle existe (décision du 03/10/2026, option A)."""
    from datetime import date

    from orisflow_engine import carnet
    from orisflow_engine.classification import classer_fichiers

    dossier_carnet = str(tmp_path / "carnet")
    carnet.enregistrer(os.path.join(dossier_carnet, carnet.NOM_FICHIER), date(2026, 10, 2), {"afriland:65": 134_381_673})

    resultat = classer_fichiers([], dossier_carnet=dossier_carnet, jour=date(2026, 10, 3))

    manquants = {r["cle"]: r for r in resultat["releves_manquants"]}
    assert len(manquants) == 10
    assert manquants["afriland:65"]["veille"] == 134_381_673
    assert manquants["cca:12"]["veille"] is None


# --- Robustesse : aucun fichier ne doit jamais arrêter le classement du lot (06/10/2026) --


def test_pdf_corrompu_nempeche_pas_le_classement_du_reste_du_lot(tmp_path):
    """Un fichier renommé en .pdf sans en être un (cas réel : téléchargement interrompu)
    ne doit jamais faire planter l'analyse des autres fichiers du même lot."""
    corrompu = tmp_path / "EtBalance_General_Consolide_corrompu.pdf"
    corrompu.write_bytes(b"ceci n'est pas un PDF valide")
    bon = tmp_path / "Akwa_Compte.xlsx"
    _extraction_comptes(bon, ["37110"])

    resultat = classer_fichiers([str(corrompu), str(bon)])

    fichier_corrompu = resultat["fichiers"][0]
    assert fichier_corrompu["niveau"] == "bloquant"
    assert fichier_corrompu["type_detecte"] == "illisible"
    fichier_bon = resultat["fichiers"][1]
    assert fichier_bon["type_detecte"] == "compte"
    assert fichier_bon["total_comptes"] == 1


def test_classeur_excel_tronque_est_signale_sans_planter(tmp_path):
    """Un classeur `.xlsx` tronqué (téléchargement interrompu) lève une exception chez
    openpyxl : le filet de sécurité générique doit la transformer en résultat bloquant."""
    tronque = tmp_path / "Mokolo_Compte.xlsx"
    tronque.write_bytes(b"PK\x03\x04 ceci n'est pas un vrai classeur Excel")

    resultat = classer_fichiers([str(tronque)])

    f = resultat["fichiers"][0]
    assert f["niveau"] == "bloquant"
    assert "n'a pas pu être lu" in f["messages"][0]


# --- Robustesse : fichier ouvert dans Excel au moment de la lecture (06/10/2026) ---------


def test_fichier_verrouille_par_excel_donne_un_message_clair(tmp_path, monkeypatch):
    """Simule le cas le plus fréquent en pratique : le fichier est encore ouvert dans Excel
    quand Orisflow essaie de le lire (PermissionError). Le message doit être compréhensible
    par un comptable, pas un nom de classe Python, et ne doit jamais planter le lot."""
    import orisflow_engine.classification as classification_module

    chemin = tmp_path / "Akwa_Compte.xlsx"
    _extraction_comptes(chemin, ["37110"])

    def leve_permission_denied(*args, **kwargs):
        raise PermissionError(13, "Permission denied", str(chemin))

    monkeypatch.setattr(classification_module, "analyser_comptes", leve_permission_denied)

    resultat = classer_fichiers([str(chemin)])
    f = resultat["fichiers"][0]
    assert f["niveau"] == "bloquant"
    assert "ouvert dans Excel" in f["messages"][-1]
    assert "PermissionError" not in f["messages"][-1]
