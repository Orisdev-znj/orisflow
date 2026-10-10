"""Tests du recalcul d'un classeur via LibreOffice (avec cache)."""

import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from orisflow_engine import recalcul
from orisflow_engine.recalcul import recalculer_classeur


def _classeur_avec_formule_non_calculee(chemin):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille["A1"] = 2
    feuille["A2"] = 3
    feuille["A3"] = "=A1+A2"  # openpyxl n'évalue jamais cette formule à l'écriture
    classeur.save(chemin)


@pytest.mark.libreoffice
def test_recalcule_une_formule_non_evaluee_puis_reutilise_le_cache(tmp_path, monkeypatch):
    chemin = tmp_path / "test.xlsx"
    _classeur_avec_formule_non_calculee(chemin)
    assert openpyxl.load_workbook(chemin, data_only=True).active["A3"].value is None

    chemin_recalcule = recalculer_classeur(str(chemin))
    assert chemin_recalcule is not None
    assert openpyxl.load_workbook(chemin_recalcule, data_only=True).active["A3"].value == 5

    # Deuxième appel sur le même fichier inchangé : servi par le cache, sans relancer LibreOffice.
    monkeypatch.setattr(recalcul, "_trouver_soffice", lambda: None)
    assert recalculer_classeur(str(chemin)) == chemin_recalcule


def test_ne_modifie_jamais_le_fichier_dorigine(tmp_path):
    chemin = tmp_path / "test.xlsx"
    _classeur_avec_formule_non_calculee(chemin)
    empreinte_avant = chemin.read_bytes()

    recalculer_classeur(str(chemin))

    assert chemin.read_bytes() == empreinte_avant


def test_fichier_introuvable_renvoie_none():
    assert recalculer_classeur(r"C:\chemin\qui\nexiste\pas.xlsx") is None


def test_soffice_absent_renvoie_none_sans_exception(tmp_path):
    chemin = tmp_path / "test.xlsx"
    _classeur_avec_formule_non_calculee(chemin)

    assert recalculer_classeur(str(chemin)) is None  # LibreOffice désactivé par conftest.py
